"""
The audit trail of one survey as an Excel workbook: what it asked, who it went
through, every step it took, and every time it was sent back with which
comments on which questions.

One sheet, "Audit report", with three stacked blocks on a shared grid of eight
columns: survey details, a timeline of every step, and one row per question
with the answer from each submission and the comments made on it.
"""
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.core.question_types import is_empty
from apps.surveys.models import Survey, SurveyQuestion

ACTION_LABELS = {
    "created": "Created",
    "published": "Published and assigned",
    "submitted": "Response submitted",
    "reviewed": "Passed review",
    "approved": "Approved",
    "returned": "Returned",
    "closed": "Closed",
}
AUDIENCE_LABELS = {"responder": "Responder", "reviewer": "Reviewer", "both": "Responder and reviewer"}
HEADER_FILL = PatternFill("solid", fgColor="1F6F78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
BAND_FILL = PatternFill("solid", fgColor="E3F0F1")
BAND_FONT = Font(bold=True, color="164F56")
LABEL_FILL = PatternFill("solid", fgColor="F5F7F7")
CHANGED_FILL = PatternFill("solid", fgColor="FFF4CF")
RETURNED_FONT = Font(bold=True, color="B4492D")
OK_FONT = Font(bold=True, color="2C7A4B")
EDGE = Side(style="thin", color="D9DFE0")
BOX = Border(left=EDGE, right=EDGE, top=EDGE, bottom=EDGE)
ZEBRA_FILL = PatternFill("solid", fgColor="F8FAFA")
WIDTHS = [8, 20, 46, 22, 24, 24, 24, 60]
WRAP = Alignment(wrap_text=True, vertical="top")


def _when(value):
    if value is None:
        return ""
    return timezone.localtime(value).strftime("%Y-%m-%d %H:%M")


def _nice(value):
    """A readable moment: 17 Oct 2026, 13:19."""
    if value is None:
        return ""
    return timezone.localtime(value).strftime("%d %b %Y, %H:%M").lstrip("0")


TYPE_LABELS = {
    "text": "Free text", "single_choice": "Single choice", "multi_choice": "Multiple choice",
    "rating": "Rating (1-5)", "yes_no": "Yes / No", "number": "Number", "date": "Date",
}


def _name(person):
    return (person.full_name or person.email) if person else ""


def _text(value):
    """An answer as a spreadsheet cell: lists joined, empties as a dash."""
    if is_empty(value):
        return "—"
    if isinstance(value, list):
        return ", ".join(str(v) for v in value)
    return value if isinstance(value, (int, float)) else str(value)


SHEET = "Audit report"


class _Sheet:
    """Writes the stacked blocks, tracking the next free row."""

    def __init__(self):
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = SHEET
        self.row = 1
        for i, width in enumerate(WIDTHS, start=1):
            self.ws.column_dimensions[get_column_letter(i)].width = width

    def band(self, title):
        ws = self.ws
        ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=8)
        cell = ws.cell(self.row, 1, title.upper())
        cell.fill, cell.font = BAND_FILL, BAND_FONT
        self.row += 1

    def header(self, labels):
        for col, label in enumerate(labels, start=1):
            cell = self.ws.cell(self.row, col, label)
            cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, WRAP
        self.row += 1

    def line(self, values):
        for col, value in enumerate(values, start=1):
            self.ws.cell(self.row, col).value = value
            self.ws.cell(self.row, col).alignment = WRAP
        self.row += 1
        return self.row - 1

    def pair(self, left, right):
        """A label/value pair on each half: labels in A:B and D:E, values in C and F:H."""
        ws, r = self.ws, self.row
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=8)
        for col, value in ((1, left[0]), (3, left[1]), (4, right[0] if right else ""), (6, right[1] if right else "")):
            ws.cell(r, col, value).alignment = WRAP
        for col in (1, 4):
            ws.cell(r, col).fill, ws.cell(r, col).font = LABEL_FILL, Font(bold=True)
        self.row += 1

    def gap(self):
        self.row += 1


def _comment_parts(comment, event, role_of):
    """One comment as rich text: a bold line (when, who, to whom), then the words."""
    to = AUDIENCE_LABELS.get(event.audience, "")
    who = f"{_name(comment.author)} ({role_of.get(comment.author.pk, 'Administrator')})" if comment.author else "Unknown"
    return [TextBlock(InlineFont(b=True), f"{_nice(event.created_at)} · {who} → {to}\n"), comment.text]


def _span(ws, row, start):
    """Stretch the cell at `start` across to column H, keeping the box border on every part."""
    if start < 8:
        ws.merge_cells(start_row=row, start_column=start, end_row=row, end_column=8)
    for col in range(start, 9):
        ws.cell(row, col).border = BOX


def _polish(ws):
    """Readability pass: no gridlines, light cell borders, zebra stripes, centred
    short columns, taller bands and headers, and a landscape page that fits one width."""
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    stripe = False
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        first = row[0]
        if first.value is None and all(c.value is None for c in row):
            stripe = False
            continue
        is_band = first.fill == BAND_FILL
        is_header = first.fill == HEADER_FILL
        if is_band:
            ws.row_dimensions[first.row].height = 24
            for c in row:
                c.alignment = Alignment(vertical="center", indent=1)
            continue
        if is_header:
            ws.row_dimensions[first.row].height = 32
            stripe = False
        else:
            stripe = not stripe
        for c in row:
            if c.fill == PatternFill():  # untouched cell
                if stripe and not is_header:
                    c.fill = ZEBRA_FILL
            if c.value is not None or c.fill != PatternFill():
                c.border = BOX
            if c.column == 1 and isinstance(c.value, int) and not is_header:
                c.alignment = Alignment(horizontal="center", vertical="top")
            if is_header:
                c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def build_workbook(survey: Survey) -> bytes:
    events = list(survey.events.select_related("by").prefetch_related("question_comments__author"))
    questions = list(
        SurveyQuestion.objects.filter(section__survey=survey).select_related("section").order_by("section__order", "order")
    )
    role_of = {}
    for role in ("responder", "reviewer", "approver"):
        person = getattr(survey, role)
        if person:
            role_of[person.pk] = role.capitalize()
    submissions = [e for e in events if e.action == "submitted"]
    returns = [e for e in events if e.action == "returned"]
    published = next((e for e in events if e.action == "published"), None)
    finished = next((e for e in reversed(events) if e.action in ("approved", "closed")), None)

    sheet = _Sheet()
    ws = sheet.ws

    # ---- 1. Survey details
    sheet.band("1 · Survey details")
    days = (finished.created_at - published.created_at).days if finished and published else ""
    left = [
        ("Survey", survey.title), ("Category", survey.category.label), ("Aimed at (org block)", survey.target_unit.name),
        ("Status", dict(Survey.STATUSES).get(survey.status, survey.status)), ("Due date", survey.due_date.isoformat() if survey.due_date else ""),
        ("Number of questions", len(questions)), ("Times submitted", len(submissions)),
        ("Question comments in total", sum(len(e.question_comments.all()) for e in returns)), ("Created by", _name(survey.created_by)),
    ]
    right = [
        ("Responder", _name(survey.responder)), ("Reviewer", _name(survey.reviewer)), ("Approver", _name(survey.approver)),
        ("Published", _when(survey.published_at)), ("Finished", _when(finished.created_at) if finished else ""),
        ("Days from publish to finish", days), ("Times returned", len(returns)),
        ("Returned to responder / reviewer", f"{sum(1 for e in returns if e.audience in ('responder', 'both'))} / {sum(1 for e in returns if e.audience in ('reviewer', 'both'))}"),
        ("Report generated", _when(timezone.now())),
    ]
    for l, r in zip(left, right):
        sheet.pair(l, r)
    for label, value in (("Description", survey.description), ("Instructions", survey.instructions)):
        if value:
            sheet.pair((label, value), None)
    sheet.gap()

    # ---- 2. Timeline
    sheet.band("2 · Timeline, every step in order")
    sheet.header(["Step", "Date and time", "Action", "By", "Role", "Returned to", "Question comments", "Comment"])
    for step, e in enumerate(events, start=1):
        r = sheet.line([
            step, _nice(e.created_at), ACTION_LABELS.get(e.action, e.action), _name(e.by),
            role_of.get(e.by_id, "Administrator" if e.by_id else ""),
            AUDIENCE_LABELS.get(e.audience, "") if e.action == "returned" else "",
            len(e.question_comments.all()) or "", e.comment,
        ])
        if e.action in ("returned", "closed"):
            ws.cell(r, 3).font = RETURNED_FONT
        elif e.action == "approved":
            ws.cell(r, 3).font = OK_FONT
    sheet.gap()

    # ---- 3. Questions: answer per submission, comments per question
    sheet.band("3 · Questions, answers by submission, and comments")
    # Only submissions whose answers were recorded get a column. Earlier ones (from
    # before answer history existed) stay in the timeline. The first three fill E:G;
    # any more continue to the right of the Comments column.
    recorded = [(n, e) for n, e in enumerate(submissions, start=1) if e.snapshot]
    sub_col = lambda i: 5 + i if i < 3 else 8 + (i - 2)
    # Comments start in the first column no submission uses and stretch to H, so
    # a survey with one recorded submission has no empty columns in between.
    first_comment = 5 + min(len(recorded), 3)
    sheet.header(["Q#", "Section", "Question", "Type"])
    header_row = sheet.row - 1
    cell = ws.cell(header_row, first_comment, "Comments on this question")
    cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, WRAP
    span = lambda r: _span(ws, r, first_comment)
    span(header_row)
    for i, (n, e) in enumerate(recorded):
        col = sub_col(i)
        cell = ws.cell(header_row, col, f"Submission {n}\n{_nice(e.created_at)}")
        cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, WRAP
        ws.column_dimensions[get_column_letter(col)].width = 24
    if len(recorded) < len(submissions):
        skipped = len(submissions) - len(recorded)
        ws.merge_cells(start_row=sheet.row, start_column=3, end_row=sheet.row, end_column=8)
        sheet.line(["", "", f"{skipped} earlier submission{'s' if skipped != 1 else ''} made before answers were recorded: see the timeline above."])
        ws.cell(sheet.row - 1, 3).font = Font(italic=True, color="5B6B6E")

    notes_by_question = {}
    for e in returns:
        for c in e.question_comments.all():
            notes_by_question.setdefault(c.question_id, []).append(_comment_parts(c, e, role_of))
    for n, q in enumerate(questions, start=1):
        r = sheet.line([n, q.section.title, q.text, TYPE_LABELS.get(q.type, q.type)])
        notes = notes_by_question.get(q.pk)
        target = ws.cell(r, first_comment)
        target.alignment = WRAP
        if notes:
            parts = []
            for k, note in enumerate(notes):
                parts += (["\n\n"] if k else []) + note
            target.value = CellRichText(*parts)
        else:
            target.value = "No comments"
            target.font = Font(color="5B6B6E")
        span(r)
        # Merged cells don't auto-fit their height, so size the row for its longest cell.
        width = sum(WIDTHS[first_comment - 1:])
        lines = sum(-(-len(t) // width) + 1 for t in (note[-1] for note in notes)) if notes else 1
        ws.row_dimensions[r].height = max(15 * max(lines + len(notes or []) - 1, -(-len(q.text) // WIDTHS[2])), 18)
        previous = None
        for i, (_, e) in enumerate(recorded):
            value = _text(e.snapshot.get(str(q.id)))
            cell = ws.cell(r, sub_col(i), value)
            cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
            if previous is not None and value != previous:
                cell.fill = CHANGED_FILL
            previous = value

    _polish(ws)
    out = BytesIO()
    sheet.wb.save(out)
    return out.getvalue()
