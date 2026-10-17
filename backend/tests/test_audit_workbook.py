"""The audit-trail workbook: sheets, return counts and per-question comments."""
from datetime import datetime, timezone
from io import BytesIO
from types import SimpleNamespace as NS

from openpyxl import load_workbook

from apps.reports import audit


class FakeQS(list):
    def select_related(self, *a):
        return self

    def prefetch_related(self, *a):
        return self

    def filter(self, **kw):
        return self

    def order_by(self, *a):
        return self


class Many:
    def __init__(self, items):
        self._items = items

    def all(self):
        return self._items


def at(day):
    return datetime(2026, 10, day, 9, 0, tzinfo=timezone.utc)


def test_workbook_describes_the_whole_flow(monkeypatch):
    section = NS(title="Processes")
    q1 = NS(pk="q1", id="q1", section=section, text="Documented?", type="yes_no", options=[], required=True, hint="")
    q2 = NS(pk="q2", id="q2", section=section, text="Rate it", type="rating", options=[], required=True, hint="")
    monkeypatch.setattr(audit.SurveyQuestion, "objects", NS(filter=lambda **kw: FakeQS([q1, q2])))

    people = {k: NS(pk=k, full_name=k.title(), email=f"{k}@x") for k in ("responder", "reviewer", "approver")}
    note = NS(question_id="q2", text="Too high, explain", author=people["approver"])
    events = [
        NS(action="created", by=None, by_id=None, created_at=at(1), comment="", audience="", snapshot={}, question_comments=Many([])),
        NS(action="submitted", by=people["responder"], by_id="responder", created_at=at(2), comment="", audience="", snapshot={"q1": "Yes", "q2": 5}, question_comments=Many([])),
        NS(action="returned", by=people["approver"], by_id="approver", created_at=at(3), comment="See Q2", audience="both", snapshot={}, question_comments=Many([note])),
        NS(action="submitted", by=people["responder"], by_id="responder", created_at=at(4), comment="", audience="", snapshot={"q1": "Yes", "q2": 3}, question_comments=Many([])),
        NS(action="approved", by=people["approver"], by_id="approver", created_at=at(5), comment="", audience="", snapshot={}, question_comments=Many([])),
    ]
    survey = NS(
        title="Pulse", description="", instructions="", status="approved", due_date=None, created_at=at(1), published_at=at(1),
        category=NS(label="Ops"), target_unit=NS(name="Eng"), created_by=None, answers={"q1": "Yes", "q2": 3},
        events=NS(select_related=lambda *a: NS(prefetch_related=lambda *b: FakeQS(events))), **people,
    )
    wb = load_workbook(BytesIO(audit.build_workbook(survey)))

    assert wb.sheetnames == ["Audit report"]
    rows = [r for r in wb["Audit report"].iter_rows(values_only=True)]
    cells = {c for r in rows for c in r if c is not None}

    # Block 1: details and counts, label then value.
    details = {r[0]: r[2] for r in rows if r[0]} | {r[3]: r[5] for r in rows if r[3]}
    assert details["Times submitted"] == 2 and details["Times returned"] == 1
    assert details["Returned to responder / reviewer"] == "1 / 1" and details["Question comments in total"] == 1

    # Block 2: five steps; the return names who it went to and how many comments it carried.
    start = next(i for i, r in enumerate(rows) if r[0] == "Step")
    timeline = rows[start + 1:start + 6]
    assert [r[2] for r in timeline] == ["Created", "Response submitted", "Returned", "Response submitted", "Approved"]
    assert timeline[2][5] == "Responder and reviewer" and timeline[2][6] == 1 and timeline[2][7] == "See Q2"

    # Block 3: one row per question, answer per submission, comments on the question.
    start = next(i for i, r in enumerate(rows) if r[0] == "Q#")
    q1_row, q2_row = rows[start + 1], rows[start + 2]
    assert q1_row[7] == "None" and q1_row[4:6] == ("Yes", "Yes")
    assert q2_row[4:6] == (5, 3)
    assert "Approver) to Responder and reviewer" in q2_row[7] and "Too high, explain" in q2_row[7]
    assert any(str(c).startswith("Submission 2") for c in cells)

    # An answer that changed since the previous submission is highlighted.
    ws = wb["Audit report"]
    assert ws.cell(start + 3, 6).fill.fgColor.rgb.endswith("FFF4CF") and ws.cell(start + 2, 6).fill.fgColor.rgb.endswith("FFF4CF") is False
